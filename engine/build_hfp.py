"""Write the workbook Owl reads for one scenario. Called with Owl's Python.

The spec is the scenario with its household's people merged in (perch.with_household).
"""

import json
import sys
from openpyxl import Workbook

spec = json.loads(open(sys.argv[1]).read())
out = sys.argv[2]
PLAN = int(spec["planStart"])
START = int(spec.get("startYear") or PLAN)
PRICE = float(spec.get("priceFactor") or 1.0)
# A plan restarted in a later year runs in today's dollars, because Owl's tax brackets are this
# year's. Amounts fixed in dollars of their own year (the mortgage, a property's cost basis) shrink
# by the inflation that has happened since. A property's value does not: Owl grows it with prices.
DEFLATE = float(spec.get("deflator") or 1.0)
HEALTH = float(spec.get("healthAbove") or 0) / 100
COLLEGE = float(spec.get("collegeAbove") or 0) / 100
PEOPLE = [p["name"] for p in spec["people"]]


def inflation() -> float:
    """The yearly inflation Owl will use for this scenario, as a decimal."""
    if not spec.get("useHistoricalRates"):
        return float(spec["inflation"]) / 100
    from owlplanner.rates import df

    window = df[(df["year"] >= 1928) & (df["year"] <= 2025)]["Inflation"] / 100
    return float((1 + window).prod() ** (1 / len(window)) - 1)


RATE = inflation()


def future(amount, year, above=0.0):
    """Today's dollars to the dollars of that year, as Owl's sheets expect, growing `above` faster than prices."""
    return abs(float(amount)) * PRICE * (1 + RATE) ** (year - START) * (1 + above) ** (year - PLAN)


def sale_year(prop):
    year = prop.get("sellYear")
    return int(year) if year else 0


spending = {}
income = {}
for child in spec.get("children") or []:
    first = int(child.get("collegeStart") or 0)
    for year in range(max(first, START), first + int(child.get("years") or 0)):
        spending[year] = spending.get(year, 0) + future(child["cost"], year, COLLEGE)
for band in spec.get("spendingBands") or []:
    for year in range(max(int(band["from"]), START), int(band["to"]) + 1):
        spending[year] = spending.get(year, 0) + future(band["amount"], year, HEALTH if band.get("grows") == "health" else 0.0)
last_year = START + 60
for prop in spec.get("properties") or []:
    if not float(prop.get("income") or 0):
        continue
    stop = sale_year(prop) or last_year
    for year in range(START, stop):
        income[year] = income.get(year, 0) + future(prop["income"], year)

# Pay and paycheck savings, per person, for each year worked.
WORK = [("anticipated wages", "wages"), ("401k ctrb", "pretaxSave"), ("Roth 401k ctrb", "rothSave"), ("HSA ctrb", "hsaSave")]


def per_person(key, i):
    values = spec.get(key) or []
    return float(values[i] or 0) if i < len(values) else 0.0


def paid_out(i, year):
    """Money that leaves W-2 pay without lowering its tax, as Owl taxes pay as entered and never
    takes savings out of cash: Roth and after-tax 401k savings, and the payroll taxes Owl does
    not charge. Entered as an untaxed outflow."""
    return round(future(per_person("rothSave", i) + per_person("payrollTax", i), year))


wb = Workbook()
for i, name in enumerate(PEOPLE):
    sheet = wb.active if i == 0 else wb.create_sheet(name)
    sheet.title = name
    until = (spec.get("workUntil") or [None] * len(PEOPLE))[i]
    worked = range(START, int(until) + 1) if until else range(0)
    # Household spending and income go on the first person's sheet; taxes are joint.
    rows = set(worked) | {START} | (set(spending) | set(income) if i == 0 else set())
    sheet.append(["year"] + [column for column, _ in WORK] + ["big-ticket items"] + (["other inc"] if i == 0 else []))
    for year in sorted(rows):
        work = [round(future(per_person(key, i), year)) if year in worked else 0 for _, key in WORK]
        # The employer match lands in the pre-tax 401k without passing through pay, and Owl
        # never takes savings out of cash, so it is simply added to the deposit.
        work[1] += round(future(per_person("matchSave", i), year)) if year in worked else 0
        leaving = paid_out(i, year) if year in worked else 0
        house = [-leaving - round(spending.get(year, 0)), round(income.get(year, 0))] if i == 0 else [-leaving]
        sheet.append([year] + work + house)

debts = wb.create_sheet("Debts")
debts.append(["active", "name", "type", "year", "term", "amount", "rate", "property"])
assets = wb.create_sheet("Fixed Assets")
assets.append(["active", "name", "type", "year", "basis", "value", "rate", "yod", "commission"])
for prop in spec.get("properties") or []:
    kind = "residence" if prop.get("use") == "live" else "real estate"
    if float(prop.get("mortgage") or 0) and int(prop.get("years") or 0):
        debts.append([True, f"{prop['name']} mortgage", "mortgage", PLAN, int(prop["years"]),
                      round(float(prop["mortgage"]) / DEFLATE), float(prop.get("rate") or 0), prop["name"]])
    assets.append([True, prop["name"], kind, PLAN, round(float(prop.get("cost") or 0) / DEFLATE),
                   round(float(prop.get("value") or 0)), float(prop.get("growth") or 0), sale_year(prop), 0])
wb.save(out)
